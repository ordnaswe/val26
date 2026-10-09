#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
distrikt_slutligt.py — röster per valdistrikt i alla tre valen ur Valmyndighetens slutliga Excel-filer
("roster-per-distrikt-slutligt-…-riksdagsvalet/regionvalen/kommunvalen-2026.xlsx") -> data/distrikt_2026.csv.gz
Långt format: valtyp, distriktskod, distrikt, kommunkod, parti, parti_abbr, roster, giltiga, rostberattigade.
Kräver openpyxl. Används av build_kommun.py (kommunsidor med distriktskarta, röstsplittring m.m.).
    python3 distrikt_slutligt.py --rd roster-…-riksdagsvalet-2026.xlsx --rf roster-…-regionvalen-2026.xlsx --kf roster-…-kommunvalen-2026-.xlsx
"""
import argparse, csv, gzip
import openpyxl
ABBR = {'Arbetarepartiet-Socialdemokraterna': 'S', 'Moderaterna': 'M', 'Sverigedemokraterna': 'SD', 'Vänsterpartiet': 'V',
        'Centerpartiet': 'C', 'Miljöpartiet de gröna': 'MP', 'Kristdemokraterna': 'KD', 'Liberalerna (tidigare Folkpartiet)': 'L'}
def main():
    ap = argparse.ArgumentParser(); ap.add_argument('--rd'); ap.add_argument('--rf'); ap.add_argument('--kf'); ap.add_argument('--out', default='data/distrikt_2026.csv.gz')
    a = ap.parse_args(); n = 0
    with gzip.open(a.out, 'wt', newline='', encoding='utf-8') as f:
        w = csv.writer(f); w.writerow(['valtyp', 'distriktskod', 'distrikt', 'kommunkod', 'parti', 'parti_abbr', 'roster', 'giltiga', 'rostberattigade'])
        for path, vt in [(a.rd, 'RD'), (a.rf, 'RF'), (a.kf, 'KF')]:
            if not path: continue
            wb = openpyxl.load_workbook(path, read_only=True)
            name = [s for s in wb.sheetnames if s.startswith('roster_')][0]; ws = wb[name]
            it = ws.iter_rows(values_only=True); head = [str(h) for h in next(it)]
            ci = {h: i for i, h in enumerate(head)}
            kod = ci.get('Valdistriktskod'); namn = ci.get('Valdistriktsnamn', ci.get('Valdistrikt')); kk = ci.get('Kommunkod')
            parti = ci['Parti/kategori']; rost = ci['Röster']; gilt = [i for h, i in ci.items() if h.startswith('Summa giltiga')][0]; rb = ci.get('Röstberättigade')
            for r in it:
                if not r or r[parti] is None: continue
                p = str(r[parti]).strip()
                if p.lower().startswith(('blank', 'ogiltig', 'övriga ogiltiga', 'ej anmäl', 'röstberättigade', 'summa', 'antal', 'valdeltagande', 'totalt')): continue
                v = int(r[rost] or 0)
                if v <= 0: continue
                w.writerow([vt, str(r[kod]), r[namn], str(r[kk]).zfill(4), p, ABBR.get(p, p), v, int(r[gilt] or 0), int(r[rb] or 0)]); n += 1
    print(f'Skrev {a.out}: {n} rader')
if __name__ == '__main__': main()
