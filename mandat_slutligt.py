#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
mandat_slutligt.py — gör data/mandat_slutligt.csv och data/valdeltagande_2026.csv ur Valmyndighetens
slutliga Excel-filer (sidan "Rådata och statistik val 2026"):
  * mandatfordelning-jamforelser-mellan-2022-och-2026.xlsx
  * roster-per-distrikt-slutligt-…-riksdagsvalet/regionvalen/kommunvalen-2026.xlsx (flikarna Valdeltagande_*)
Kräver openpyxl (pip install openpyxl). De färdiga csv-filerna committas, så skriptet behöver bara
köras om Valmyndigheten rättar filerna.

    python3 mandat_slutligt.py --mandat mandatfordelning-….xlsx --rd roster-…-riksdagsvalet-2026.xlsx \
        --rf roster-…-regionvalen-2026.xlsx --kf roster-…-kommunvalen-2026-.xlsx
"""
import argparse, csv
import openpyxl

ABBR = {'Arbetarepartiet-Socialdemokraterna': 'S', 'Moderaterna': 'M', 'Sverigedemokraterna': 'SD',
        'Vänsterpartiet': 'V', 'Centerpartiet': 'C', 'Miljöpartiet de gröna': 'MP',
        'Kristdemokraterna': 'KD', 'Liberalerna (tidigare Folkpartiet)': 'L'}

def rows(ws):
    it = ws.iter_rows(values_only=True); head = [str(h) for h in next(it)]
    for r in it:
        if r and any(v is not None for v in r): yield dict(zip(head, r))

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--mandat', required=True); ap.add_argument('--rd'); ap.add_argument('--rf'); ap.add_argument('--kf')
    ap.add_argument('--out', default='data/mandat_slutligt.csv'); ap.add_argument('--out-deltagande', default='data/valdeltagande_2026.csv')
    a = ap.parse_args()
    wb = openpyxl.load_workbook(a.mandat, read_only=True)
    out = []
    for sheet, valtyp, kodcol in [('Riksdag', 'RD', None), ('Region', 'RF', 'Regionkod'), ('Kommun', 'KF', 'Kommunkod')]:
        for r in rows(wb[sheet]):
            parti = (r.get('Parti') or '').strip()
            if not parti or parti.lower().startswith('totalt'): continue
            kod = '00' if kodcol is None else str(r[kodcol]).zfill(2 if valtyp == 'RF' else 4)
            out.append([valtyp, kod, ABBR.get(parti, parti), parti, int(r['Antal mandat 2022'] or 0), int(r['Antal mandat 2026'] or 0)])
    with open(a.out, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f); w.writerow(['valtyp', 'kod', 'parti_abbr', 'parti', 'mandat2022', 'mandat2026', 'kalla']); 
        for r in out: w.writerow(r + ['Valmyndigheten: Mandatfördelning riksdag, region- och kommunfullmäktige 2026 (xlsx)'])
    print(f'Skrev {a.out}: {len(out)} rader')
    delt = []
    for path, valtyp in [(a.rd, 'RD'), (a.rf, 'RF'), (a.kf, 'KF')]:
        if not path: continue
        wbv = openpyxl.load_workbook(path, read_only=True)
        for name in wbv.sheetnames:
            if not name.startswith('Valdeltagande_'): continue
            niva = name.split('_', 1)[1]
            for r in rows(wbv[name]):
                if niva == 'kommun': kod = str(r.get('Kommunkod')).zfill(4); namn = r.get('Kommun')
                elif niva in ('lan', 'län'): kod = str(r.get('Länskod')).zfill(2); namn = r.get('Län')
                elif niva == 'region': kod = str(r.get('Regionkod') or r.get('Länskod')).zfill(2); namn = r.get('Region') or r.get('Län')
                else: continue
                v = r.get('Valdeltagande (%)')
                if v is None: continue
                delt.append([valtyp, niva, kod, namn, r.get('Röster'), r.get('Röstberättigade'), round(float(v)*100, 2)])
    if delt:
        with open(a.out_deltagande, 'w', newline='', encoding='utf-8') as f:
            w = csv.writer(f); w.writerow(['valtyp', 'niva', 'kod', 'namn', 'roster', 'rostberattigade', 'valdeltagande_pct']); w.writerows(delt)
        print(f'Skrev {a.out_deltagande}: {len(delt)} rader')

if __name__ == '__main__':
    main()
